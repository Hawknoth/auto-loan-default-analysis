from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "paper"
FIGURE = ROOT / "reports" / "figures" / "portfolio_dashboard.png"
DOCX = OUT / "auto_loan_portfolio_analysis_APA7.docx"
MD = OUT / "auto_loan_portfolio_analysis_APA7.md"
TITLE = "Default and Risk Analysis of an Auto Loan Portfolio"


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    node = props.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        props.append(node)
    node.set(qn("w:fill"), fill)


def margins(cell, value=100):
    props = cell._tc.get_or_add_tcPr()
    box = props.first_child_found_in("w:tcMar")
    if box is None:
        box = OxmlElement("w:tcMar")
        props.append(box)
    for side in ("top", "start", "bottom", "end"):
        node = box.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            box.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, text, end])


def body(doc, text, indent=True):
    p = doc.add_paragraph(text)
    p.paragraph_format.line_spacing = 2
    p.paragraph_format.space_after = Pt(0)
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.5)
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing = 2
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(0)
    return p


def configure(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    add_page_number(section.header.paragraphs[0])
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    for name in ("Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(12)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.styles["Heading 2"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT


def centered(doc, text, bold=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 2
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    return p


def add_title_and_abstract(doc):
    for _ in range(4):
        doc.add_paragraph()
    centered(doc, TITLE, bold=True)
    for line in (
        "Adrian",
        "Professional Data Analytics Portfolio",
        "Business Intelligence Project",
        "September 20, 2026",
    ):
        centered(doc, line)
    doc.add_page_break()
    centered(doc, "Abstract", bold=True)
    abstract = (
        "This project examines 233,154 auto loans to identify portfolio segments that require closer "
        "monitoring from a loan servicing and collections perspective. The analysis focuses on five "
        "operational variables: disbursed amount, loan-to-value ratio, credit score, recent delinquency, "
        "and geographic location. The portfolio recorded an overall default rate of 21.7%. Loans with "
        "an LTV between 80% and 90% reached 25.9%, while customers with a delinquent account during the "
        "previous six months reached 27.1%. The 301-600 credit-score group also recorded 27.1%. These "
        "results support simple monitoring rules before predictive modeling is considered. The scope is "
        "descriptive because the dataset does not contain monthly payment history, days past due, promises "
        "to pay, cure rate, roll rate, recoveries, charge-off events, or total-loss events."
    )
    body(doc, abstract, indent=False)
    p = body(doc, "Keywords: auto loans, default, loan servicing, collections, business intelligence", indent=False)
    p.runs[0].italic = True
    doc.add_page_break()


def add_summary_table(doc):
    cap = doc.add_paragraph("Table 1\nSummary of the Main Portfolio Segments")
    cap.paragraph_format.keep_with_next = True
    cap.paragraph_format.space_before = Pt(6)
    cap.paragraph_format.space_after = Pt(6)
    cap.runs[0].bold = True
    cap.runs[-1].italic = True
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(2.7), Inches(1.25), Inches(1.25), Inches(1.25)]
    for cell, label, width in zip(table.rows[0].cells, ("Segment", "Loans", "Default", "Difference"), widths):
        cell.width = width
        cell.text = label
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cell, "1F4E78")
        margins(cell)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.name = "Times New Roman"
            run.font.size = Pt(10)
            run.bold = True
    rows = (
        ("Full portfolio", "233,154", "21.7%", "Baseline"),
        ("80%-90% LTV", "81,005", "25.9%", "+4.2 pp"),
        ("Recent delinquency", "18,195", "27.1%", "+5.4 pp"),
        ("301-600 credit score", "18,716", "27.1%", "+5.4 pp"),
        ("751-900 credit score", "25,682", "14.7%", "-7.0 pp"),
        ("State 13", "17,884", "30.7%", "+9.0 pp"),
    )
    for row_number, values in enumerate(rows, 1):
        cells = table.add_row().cells
        for index, (cell, value, width) in enumerate(zip(cells, values, widths)):
            cell.width = width
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins(cell)
            if row_number % 2 == 0:
                shade(cell, "EAF2F8")
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            for run in cell.paragraphs[0].runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(10)


def build_docx():
    OUT.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure(doc)
    add_title_and_abstract(doc)
    centered(doc, TITLE, bold=True)
    body(doc, (
        "The purpose of this project was to turn an auto loan dataset into information that a servicing or "
        "collections team could interpret. I did not attempt to build a complex prediction model. The project "
        "demonstrates a reproducible business intelligence process: clean the data, define understandable "
        "segments, calculate portfolio metrics, test differences, and translate the results into operational recommendations."
    ))
    body(doc, (
        "The main question was which borrower and loan characteristics separate accounts with an above-average "
        "default rate. The analysis therefore considered loan-to-value ratio (LTV), credit score, recent "
        "delinquency, employment type, and geographic distribution."
    ))

    heading(doc, "Method")
    heading(doc, "Data Source", 2)
    body(doc, (
        "The analysis uses the Vehicle Loan Default Prediction dataset published on Kaggle (Paul, n.d.). The "
        "training file contains 233,154 records and 41 fields. Each record represents one loan, and LOAN_DEFAULT "
        "indicates whether the account ended in default. The repository does not redistribute the original file. "
        "It provides manual download instructions and retains only aggregated outputs created by the analysis."
    ))
    heading(doc, "Data Preparation and Analysis", 2)
    body(doc, (
        "Column names were standardized, while LTV and credit score were grouped into ranges that a business "
        "reader could interpret. Recent delinquency was converted into a binary indicator based on whether the "
        "customer had a delinquent account during the previous six months. Loan volume, disbursed amount, average "
        "LTV, and default rate were calculated for each segment. Python produced the analysis and dashboard, and "
        "SQL queries replicated the principal portfolio metrics."
    ))

    heading(doc, "Results")
    heading(doc, "Portfolio Overview", 2)
    body(doc, (
        "The portfolio contains 233,154 loans and 12.67 billion in total disbursed amount, measured in the "
        "dataset's unspecified currency. Average LTV is 74.7%, and the overall default rate is 21.7%. This rate "
        "serves as the baseline for every segment comparison."
    ))
    add_summary_table(doc)
    heading(doc, "Loan to Value Ratio", 2)
    body(doc, (
        "Default increased from 13.9% among loans at or below 60% LTV to 25.9% in the 80%-90% band. The segment "
        "above 90% recorded 20.6%, but it contains only 883 loans. I therefore did not treat the relationship as "
        "perfectly linear. The operationally useful finding is that the 80%-90% band combines an elevated default "
        "rate with 81,005 accounts, creating greater exposure than the much smaller group above 90%."
    ))
    heading(doc, "Credit Score", 2)
    body(doc, (
        "Customers scoring 301-600 recorded the highest default rate at 27.1%. The 1-300 group reached 25.2%, "
        "and borrowers without a score reached 23.1%. Customers scoring 751-900 had the lowest rate at 14.7%. "
        "The no-score group should not be treated only as missing data because it contains 116,950 loans and "
        "performs worse than the portfolio baseline."
    ))
    heading(doc, "Recent Delinquency", 2)
    body(doc, (
        "Customers with at least one delinquent account during the previous six months recorded a 27.1% default "
        "rate. Customers without recent delinquency recorded 21.3%. The 5.8 percentage-point difference supports "
        "the use of recent delinquency as a straightforward early-monitoring criterion."
    ))
    heading(doc, "Geographic Distribution", 2)
    body(doc, (
        "Among states with at least 1,000 loans, State 13 had the highest default rate at 30.7% across 17,884 "
        "accounts. This result does not show that location causes default. It identifies a region that combines "
        "meaningful volume with above-average risk and therefore deserves further review before collections "
        "capacity is allocated."
    ))

    if FIGURE.exists():
        doc.add_page_break()
        p = doc.add_paragraph("Figure 1")
        p.runs[0].bold = True
        p.paragraph_format.space_after = Pt(0)
        p = doc.add_paragraph("Auto Loan Portfolio Default and Risk Overview")
        p.runs[0].italic = True
        p.paragraph_format.space_after = Pt(6)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(FIGURE), width=Inches(6.45))
        p = doc.add_paragraph("Note. Created from the Vehicle Loan Default Prediction dataset.")
        p.runs[0].italic = True
        p.paragraph_format.space_before = Pt(4)

    heading(doc, "Discussion")
    body(doc, (
        "The findings support an understandable monitoring rule without presenting the project as an automated "
        "decision system. An early-review queue could prioritize accounts that combine 80%-90% LTV, recent "
        "delinquency, and low or mid-range credit scores. State 13 also merits review because of its volume and "
        "default rate. These conditions can organize analysis, but they are not sufficient grounds for an adverse "
        "action against a customer."
    ))
    body(doc, (
        "A production collections analysis would add monthly payment and contact history. Those fields would make "
        "it possible to calculate roll rates, cure rate, promise-to-pay results, contact rate, recoveries, and "
        "movement toward charge-off. They would also help separate an origination problem from a servicing or "
        "collections strategy problem."
    ))

    heading(doc, "Limitations")
    body(doc, (
        "The dataset provides a binary default outcome but does not show the sequence that led to it. It does not "
        "contain days past due, partial payments, customer contacts, repossession, total loss, recovered balances, "
        "or charge-off dates. It also does not identify the currency or provide state names. The findings should "
        "therefore be interpreted as associations within this portfolio, not causal relationships or complete "
        "collections performance measures."
    ))

    heading(doc, "Conclusion")
    body(doc, (
        "The analysis identified three practical monitoring signals: 80%-90% LTV, recent delinquency, and a "
        "301-600 credit score. Each group performs worse than the portfolio baseline and can support initial "
        "account prioritization. The results also show why volume matters. A small segment with a high rate does "
        "not necessarily represent the largest exposure for the collections team."
    ))
    body(doc, (
        "The final repository matches the intended scope of an intermediate business intelligence project. It "
        "documents the data source, preparation steps, SQL queries, aggregated tables, dashboard, findings, and "
        "limitations. The project can be explained in an interview without relying on a model that is difficult "
        "to defend."
    ))

    doc.add_page_break()
    heading(doc, "References")
    p = body(doc, (
        "Paul, A. (n.d.). Vehicle loan default prediction [Data set]. Kaggle. "
        "https://www.kaggle.com/datasets/avikpaul4u/vehicle-loan-default-prediction"
    ), indent=False)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)
    doc.core_properties.title = TITLE
    doc.core_properties.author = "Adrian"
    doc.core_properties.subject = "Auto loan portfolio default and risk analysis"
    doc.core_properties.keywords = "auto loans, default, loan servicing, collections, business intelligence"
    doc.save(DOCX)


def build_markdown():
    content = f"""# {TITLE}

**Author:** Adrian  
**Affiliation:** Professional Data Analytics Portfolio  
**Date:** September 20, 2026

## Abstract

This project examines 233,154 auto loans to identify portfolio segments that require closer monitoring from a loan servicing and collections perspective. The overall default rate was 21.7%. Loans with an LTV between 80% and 90% reached 25.9%, customers with recent delinquency reached 27.1%, and the 301-600 credit-score group also recorded 27.1%. The analysis is descriptive because the dataset does not contain monthly payment history, days past due, cure rate, roll rate, charge-off events, or total-loss events.

*Keywords: auto loans, default, loan servicing, collections, business intelligence*

## Method

The analysis uses the *Vehicle Loan Default Prediction* dataset published on Kaggle (Paul, n.d.). The file contains 233,154 records and 41 fields. Columns were standardized and new segments were created for LTV, credit score, and recent delinquency. Python and SQL were used to calculate loan volume, disbursed amount, average LTV, and default rate.

## Results

| Segment | Loans | Default rate | Difference from portfolio |
|---|---:|---:|---:|
| Full portfolio | 233,154 | 21.7% | Baseline |
| 80%-90% LTV | 81,005 | 25.9% | +4.2 pp |
| Recent delinquency | 18,195 | 27.1% | +5.4 pp |
| 301-600 credit score | 18,716 | 27.1% | +5.4 pp |
| 751-900 credit score | 25,682 | 14.7% | -7.0 pp |
| State 13 | 17,884 | 30.7% | +9.0 pp |

### Loan to Value Ratio

Default increased from 13.9% among loans at or below 60% LTV to 25.9% in the 80%-90% band. The segment above 90% recorded 20.6%, but it contains only 883 loans. The 80%-90% band is more operationally relevant because it combines an elevated rate with 81,005 accounts.

### Credit Score

Customers scoring 301-600 recorded 27.1% default, compared with 14.7% among customers scoring 751-900. The no-score group contains 116,950 loans and recorded 23.1%, so it should not be treated only as missing data.

### Recent Delinquency

Customers with at least one delinquent account during the previous six months recorded 27.1% default, compared with 21.3% among customers without recent delinquency.

### Geographic Distribution

Among states with at least 1,000 loans, State 13 recorded the highest default rate at 30.7% across 17,884 accounts. This does not establish causation, but it identifies a concentration that deserves further review.

![Portfolio dashboard](../reports/figures/portfolio_dashboard.png)

## Discussion

An early-review queue could consider accounts that combine 80%-90% LTV, recent delinquency, and low or mid-range credit scores. A production analysis would require monthly payment and contact history to calculate roll rates, cure rate, promises to pay, recoveries, and movement toward charge-off.

## Limitations

The dataset does not include days past due, partial payments, customer contacts, repossession, total loss, recovered balances, or charge-off dates. It also does not identify the currency or provide state names. The findings represent associations within this portfolio, not causal relationships.

## Conclusion

The project identifies three practical monitoring signals and documents them at an appropriate level for an intermediate business intelligence portfolio. The repository includes the data source, processing steps, SQL queries, aggregated tables, dashboard, findings, and limitations.

## References

Paul, A. (n.d.). *Vehicle loan default prediction* [Data set]. Kaggle. https://www.kaggle.com/datasets/avikpaul4u/vehicle-loan-default-prediction
"""
    MD.write_text(content, encoding="utf-8")


if __name__ == "__main__":
    build_docx()
    build_markdown()
    print(DOCX)
    print(MD)
